# -*- coding: utf-8 -*-
"""小易 XLSX 导出单测：真实文件 + 内存仓储，不连接 MySQL。"""

import copy
import hashlib
import json
import os
import tempfile
import unittest
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from common.kb_qa.export_schema import (
    query_export_ddl_statements,
    query_snapshot_ddl_statements,
)
from common.kb_qa.query_export import (
    ExportLimitError,
    ExportRepository,
    ExportService,
    QuerySnapshotRepository,
)


class _Cursor:
    def __init__(self, *, one=None, many=(), rowcount=1):
        self.one = one
        self.many = list(many)
        self.rowcount = rowcount
        self.executed = []
        self.closed = False

    def execute(self, sql, params=None):
        self.executed.append((sql, params))

    def fetchone(self):
        return self.one

    def fetchall(self):
        return list(self.many)

    def close(self):
        self.closed = True


class _Connection:
    def __init__(self, cursor=None):
        self.cursor_instance = cursor or _Cursor()
        self.commits = 0
        self.rollbacks = 0

    def cursor(self):
        return self.cursor_instance

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1

    def close(self):
        pass


class _MemorySnapshots:
    def __init__(self):
        self.rows = {}

    def put(self, snapshot_id, sender_uid, conversation_id, payload):
        self.rows[snapshot_id] = {
            "snapshot_id": snapshot_id,
            "sender_uid": sender_uid,
            "conversation_id": conversation_id,
            "payload": copy.deepcopy(payload),
        }

    def get_for_owner(self, snapshot_id, sender_uid, conversation_id):
        row = self.rows.get(snapshot_id)
        if not row or row["sender_uid"] != sender_uid or row["conversation_id"] != conversation_id:
            return None
        return copy.deepcopy(row["payload"])

    def latest_id_for_owner(self, sender_uid, conversation_id):
        rows = [
            row for row in self.rows.values()
            if row["sender_uid"] == sender_uid and row["conversation_id"] == conversation_id
        ]
        return rows[-1]["snapshot_id"] if rows else None


class _MemoryExports:
    def __init__(self):
        self.rows = {}

    def create(self, **values):
        summary = values.pop("query_summary")
        row = dict(values)
        row.update({
            "status": "pending",
            "query_summary_json": json.dumps(summary, ensure_ascii=False, default=str),
            "row_count": None,
            "file_size": None,
            "file_sha256": None,
            "error": None,
            "finished_at": None,
            "downloaded_at": None,
            "sent_at": None,
        })
        self.rows[row["export_id"]] = row
        return row["export_id"]

    def get(self, export_id):
        row = self.rows.get(export_id)
        return copy.deepcopy(row) if row else None

    def list_expired(self, now):
        return [
            copy.deepcopy(row) for row in self.rows.values()
            if row["expires_at"] <= now
            and row["status"] in ("pending", "running", "completed", "failed")
        ]

    def claim_pending(self, now, *, stale_seconds=300):
        for task in self.rows.values():
            started = task.get("started_at")
            claimable = task["status"] == "pending" or (
                task["status"] == "running" and started
                and started <= now - timedelta(seconds=stale_seconds)
            )
            if task["expires_at"] > now and claimable:
                task["status"] = "running"
                task["started_at"] = now
                return task["export_id"]
        return None

    def mark_running(self, export_id):
        row = self.rows[export_id]
        if row["status"] != "pending":
            return False
        row["status"] = "running"
        row["error"] = None
        return True

    def mark_completed(self, export_id, *, stored_name, row_count, file_size,
                       file_sha256, finished_at):
        row = self.rows[export_id]
        if row["status"] != "running":
            return False
        row.update({
            "status": "completed", "stored_name": stored_name,
            "row_count": row_count, "file_size": file_size,
            "file_sha256": file_sha256, "finished_at": finished_at, "error": None,
        })
        return True

    def mark_failed(self, export_id, error, finished_at=None):
        row = self.rows[export_id]
        if row["status"] not in ("pending", "running"):
            return False
        row.update({"status": "failed", "error": str(error), "finished_at": finished_at})
        return True

    def mark_downloaded(self, export_id, downloaded_at=None):
        row = self.rows[export_id]
        if row["status"] != "completed":
            return False
        row["downloaded_at"] = downloaded_at
        return True

    def mark_sent(self, export_id, sent_at=None):
        row = self.rows[export_id]
        if row["status"] != "completed":
            return False
        row["sent_at"] = sent_at
        return True

    def expire(self, export_id):
        row = self.rows[export_id]
        if row["status"] not in ("pending", "running", "completed", "failed"):
            return False
        row["status"] = "expired"
        return True


class DdlAndRepositoryTests(unittest.TestCase):
    def test_independent_ddl_factories_and_narrow_task_table(self):
        snapshot_statements = query_snapshot_ddl_statements()
        export_statements = query_export_ddl_statements()
        self.assertEqual(1, len(snapshot_statements))
        self.assertEqual(1, len(export_statements))
        snapshot = snapshot_statements[0]
        export = export_statements[0]
        self.assertIn("kb_qa_query_snapshot", snapshot)
        self.assertNotIn("kb_qa_export`", snapshot)
        self.assertIn("UNIQUE KEY", snapshot)
        self.assertIn("payload_json` LONGTEXT", snapshot)
        self.assertIn("datetime(6)", snapshot)
        for field in (
            "file_sha256", "freshness_status", "query_summary_json", "expires_at",
            "downloaded_at", "sent_at", "stored_name",
        ):
            self.assertIn(f"`{field}`", export)
        self.assertNotIn("payload_json", export)
        self.assertNotIn("message_content", export)
        self.assertNotIn("full_sql", export)

    def test_snapshot_repository_serializes_decodes_and_scopes_owner(self):
        connection = _Connection()
        QuerySnapshotRepository(connection).save(
            "s" * 32, "msg-1", "u1", "c1", {"when": date(2026, 1, 2)},
        )
        sql, params = connection.cursor_instance.executed[0]
        self.assertIn("INSERT INTO `kb_qa_query_snapshot`", sql)
        self.assertEqual(json.loads(params[4]), {"when": "2026-01-02"})
        self.assertEqual(params[:4], ("s" * 32, "msg-1", "u1", "c1"))

        cursor = _Cursor(one={"payload_json": '{"version":1,"items":[]}'})
        repo = QuerySnapshotRepository(_Connection(cursor))
        self.assertEqual({"version": 1, "items": []}, repo.get_for_owner("s1", "u1", "c1"))
        self.assertEqual(cursor.executed[0][1], ("s1", "u1", "c1"))
        self.assertIn("`sender_uid` = %s", cursor.executed[0][0])

        cursor = _Cursor(one=None)
        repo = QuerySnapshotRepository(_Connection(cursor))
        self.assertIsNone(repo.latest_for_owner("u1", "c1"))
        self.assertIn("ORDER BY `created_at` DESC", cursor.executed[0][0])
        self.assertEqual(cursor.executed[0][1], ("u1", "c1"))

    def test_export_repository_transitions_are_parameterized_and_conditional(self):
        cursor = _Cursor()
        repo = ExportRepository(_Connection(cursor))
        now = datetime(2026, 1, 2, 3, 4, 5)
        repo.create(
            export_id="e" * 32, snapshot_id="s" * 32, requester_uid="u1",
            conversation_id="c1", file_name="小易.xlsx", stored_name="random.xlsx",
            freshness_status="fresh", query_summary={"items": []},
            expires_at=now, created_at=now,
        )
        self.assertTrue(repo.mark_running("e" * 32))
        self.assertTrue(repo.mark_completed(
            "e" * 32, stored_name="random.xlsx", row_count=2, file_size=100,
            file_sha256="a" * 64, finished_at=now,
        ))
        self.assertTrue(repo.mark_downloaded("e" * 32, now))
        self.assertTrue(repo.mark_sent("e" * 32, now))
        self.assertTrue(repo.mark_failed("e" * 32, "boom", now))
        self.assertTrue(repo.expire("e" * 32))
        sql_text = "\n".join(sql for sql, _params in cursor.executed)
        self.assertIn("AND `status` = 'pending'", sql_text)
        self.assertIn("AND `status` = 'running'", sql_text)
        self.assertIn("AND `status` = 'completed'", sql_text)
        self.assertIn("('pending', 'running', 'completed', 'failed')", sql_text)
        self.assertNotIn("boom", sql_text)
        self.assertTrue(all(params is not None for _sql, params in cursor.executed))

        cursor = _Cursor(many=[{"export_id": "e" * 32, "status": "completed"}])
        rows = ExportRepository(_Connection(cursor)).list_expired(now)
        self.assertEqual(1, len(rows))
        self.assertEqual((now,), cursor.executed[0][1])
        self.assertIn("`expires_at` <= %s", cursor.executed[0][0])


def _typed_payload():
    return {
        "version": 1,
        "items": [{
            "tool_name": "daily_sales",
            "effective_params": {"scope": "杭州", "date": "2026-01-02"},
            "columns": [
                {"key": "qty", "label": "数量", "unit": "件", "description": "出库数量"},
                {"key": "amount", "label": "金额", "unit": "元", "description": "含税金额"},
                {"key": "ratio", "label": "完成率"},
                {"key": "day", "label": "业务日期"},
                {"key": "moment", "label": "导出数据时间"},
                {"key": "enabled", "label": "是否启用"},
                {"key": "missing", "label": "空值"},
                {"key": "formula", "label": "公式攻击"},
                {"key": "plus", "label": "加号攻击"},
                {"key": "minus", "label": "减号攻击"},
                {"key": "at", "label": "At 攻击"},
            ],
            "rows": [{
                "qty": 0,
                "amount": Decimal("12.50"),
                "ratio": 0.0,
                "day": date(2026, 1, 2),
                "moment": datetime(2026, 1, 2, 3, 4, 5, 6000),
                "enabled": False,
                "missing": None,
                "formula": "  =1+1",
                "plus": "+CMD",
                "minus": "\t-2+3",
                "at": " @SUM(A1:A2)",
            }],
            "data_scope": "杭州区域日销口径",
            "data_version": "2026-01-02 23:59:59",
            "freshness_status": "fresh",
        }],
    }


class ExportServiceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.export_dir = Path(self.directory.name) / "exports"
        self.clock = [datetime(2026, 1, 2, 12, 0, 0)]
        self.snapshots = _MemorySnapshots()
        self.exports = _MemoryExports()
        self.service = self._service()

    def _service(self, **overrides):
        values = {
            "now": lambda: self.clock[0],
            "snapshot_repository_factory": lambda _connection: self.snapshots,
            "export_repository_factory": lambda _connection: self.exports,
        }
        values.update(overrides)
        return ExportService(lambda: _Connection(), self.export_dir, **values)

    def _snapshot(self, payload=None, snapshot_id="s" * 32):
        self.snapshots.put(snapshot_id, "u1", "c1", payload or _typed_payload())
        return snapshot_id

    def test_real_xlsx_types_formula_protection_notes_fields_and_download(self):
        snapshot_id = self._snapshot()
        export_id = self.service.create(snapshot_id, "u1", "c1")
        self.assertEqual(32, len(export_id))
        pending = self.exports.rows[export_id]
        self.assertEqual("pending", pending["status"])
        self.assertTrue(pending["file_name"].endswith(".xlsx"))
        self.assertIn("小易区域日销", pending["file_name"])
        self.assertFalse(list(self.export_dir.glob("*.xlsx")))

        path = self.service.generate(export_id)
        self.assertTrue(path.is_file())
        self.assertEqual(pending["stored_name"], path.name)
        self.assertNotEqual(pending["file_name"], path.name)
        completed = self.exports.rows[export_id]
        self.assertEqual("completed", completed["status"])
        self.assertEqual(1, completed["row_count"])
        self.assertEqual(path.stat().st_size, completed["file_size"])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), completed["file_sha256"])

        workbook = load_workbook(path, data_only=False)
        self.assertEqual(["数据", "导出说明", "字段说明"], workbook.sheetnames)
        sheet = workbook["数据"]
        self.assertEqual("数量（件）", sheet.cell(1, 1).value)
        self.assertEqual("金额（元）", sheet.cell(1, 2).value)
        self.assertEqual(0, sheet.cell(2, 1).value)
        self.assertEqual(12.5, sheet.cell(2, 2).value)
        self.assertEqual(0, sheet.cell(2, 3).value)
        # XLSX 日期与时间都是带格式的序列值，openpyxl 读取 date 单元格为午夜 datetime。
        self.assertEqual(datetime(2026, 1, 2), sheet.cell(2, 4).value)
        self.assertEqual("yyyy-mm-dd", sheet.cell(2, 4).number_format)
        self.assertEqual(datetime(2026, 1, 2, 3, 4, 5, 6000), sheet.cell(2, 5).value)
        self.assertIs(False, sheet.cell(2, 6).value)
        self.assertIsNone(sheet.cell(2, 7).value)
        for column in range(8, 12):
            cell = sheet.cell(2, column)
            self.assertTrue(cell.value.startswith("'"), (column, cell.value))
            self.assertNotEqual("f", cell.data_type)
        self.assertEqual("'  =1+1", sheet.cell(2, 8).value)
        self.assertEqual("A2", sheet.freeze_panes)
        self.assertTrue(sheet.auto_filter.ref)
        self.assertGreater(sheet.column_dimensions["A"].width, 10)
        workbook.close()

        workbook = load_workbook(path, data_only=False)
        notes = workbook["导出说明"]
        note_text = "\n".join(str(cell.value) for row in notes.iter_rows() for cell in row)
        for expected in (
            "导出人", "u1", "导出时间", "条件摘要", "杭州", "口径/数据范围",
            "来源", "数据截至", "2026-01-02 23:59:59", "行数", "freshness", "fresh",
        ):
            self.assertIn(expected, note_text)
        fields = list(workbook["字段说明"].iter_rows(values_only=True))
        self.assertIn(("数据", "qty", "数量", "件", "出库数量"), fields)
        workbook.close()

        downloaded_path, file_name = self.service.download(export_id, "u1", "c1")
        self.assertEqual(path, downloaded_path)
        self.assertEqual(pending["file_name"], file_name)
        self.assertIsNotNone(self.exports.rows[export_id]["downloaded_at"])

    def test_empty_result_and_multiple_tools_have_numbered_data_sheets(self):
        payload = {
            "version": 1,
            "items": [
                {
                    "tool_name": "daily_sales", "effective_params": {},
                    "columns": [{"key": "qty", "label": "数量"}], "rows": [],
                },
                {
                    "tool_name": "channel_sales", "effective_params": {"channel": "京东"},
                    "columns": ["channel"], "rows": [{"channel": "京东"}],
                },
            ],
        }
        export_id = self.service.create(self._snapshot(payload), "u1", "c1")
        path = self.service.generate(export_id)
        workbook = load_workbook(path)
        self.assertEqual(["数据", "数据_2", "导出说明", "字段说明"], workbook.sheetnames)
        self.assertEqual(1, workbook["数据"].max_row)
        self.assertEqual("数量", workbook["数据"]["A1"].value)
        self.assertEqual("渠道", workbook["数据_2"]["A1"].value)
        self.assertEqual("京东", workbook["数据_2"]["A2"].value)
        self.assertEqual(1, self.exports.rows[export_id]["row_count"])
        workbook.close()

    def test_row_limit_marks_failed_and_removes_all_files(self):
        payload = _typed_payload()
        payload["items"][0]["rows"] *= 2
        service = self._service(max_rows=1)
        export_id = service.create(self._snapshot(payload), "u1", "c1")
        with self.assertRaisesRegex(ExportLimitError, "row count"):
            service.generate(export_id)
        self.assertEqual("failed", self.exports.rows[export_id]["status"])
        self.assertEqual([], list(self.export_dir.iterdir()))

    def test_size_limit_marks_failed_and_removes_temporary_file(self):
        service = self._service(max_file_mb=0.000000001)
        export_id = service.create(self._snapshot(), "u1", "c1")
        with self.assertRaisesRegex(ExportLimitError, "file size"):
            service.generate(export_id)
        self.assertEqual("failed", self.exports.rows[export_id]["status"])
        self.assertEqual([], list(self.export_dir.iterdir()))

    def test_switch_ownership_tool_whitelist_and_authorizer_are_enforced(self):
        snapshot_id = self._snapshot()
        with self.assertRaises(PermissionError):
            self._service(enabled=False).create(snapshot_id, "u1", "c1")
        with self.assertRaises(PermissionError):
            self.service.create(snapshot_id, "other", "c1")
        with self.assertRaises(PermissionError):
            self.service.create(snapshot_id, "u1", "c1", authorizer=lambda *_: False)

        forbidden = _typed_payload()
        forbidden["items"][0]["tool_name"] = "search_docs"
        with self.assertRaises(PermissionError):
            self.service.create(self._snapshot(forbidden, "f" * 32), "u1", "c1")

        export_id = self.service.create(snapshot_id, "u1", "c1")
        with self.assertRaises(PermissionError):
            self.service.generate(export_id, authorizer=lambda *_: False)
        self.assertEqual("pending", self.exports.rows[export_id]["status"])
        self.service.generate(export_id)
        with self.assertRaises(PermissionError):
            self.service.download(export_id, "other", "c1")
        with self.assertRaises(PermissionError):
            self.service.download(export_id, "u1", "c1", authorizer=lambda *_: False)
        self.assertIsNone(self.exports.rows[export_id]["downloaded_at"])

    def test_complete_sql_and_message_do_not_enter_task_summary(self):
        payload = {
            "version": 1,
            "items": [{
                "tool_name": "query_sql",
                "effective_params": {"sql": "SELECT secret FROM fact_order_line", "date": "2026-01-02"},
                "columns": ["qty"], "rows": [{"qty": 1}],
            }],
            "message_content": "不要写入任务表",
        }
        export_id = self.service.create(self._snapshot(payload), "u1", "c1")
        summary = self.exports.rows[export_id]["query_summary_json"]
        self.assertNotIn("SELECT secret", summary)
        self.assertNotIn("不要写入任务表", summary)
        self.assertIn("已省略", summary)
        path = self.service.generate(export_id)
        workbook = load_workbook(path)
        notes = "\n".join(str(cell.value) for row in workbook["导出说明"] for cell in row)
        self.assertNotIn("SELECT secret", notes)
        workbook.close()

    def test_latest_snapshot_status_and_group_send_audit(self):
        snapshot_id = self._snapshot()
        self.assertEqual(snapshot_id, self.service.latest_snapshot_id("u1", "c1"))
        export_id = self.service.create(snapshot_id, "u1", "c1")
        path = self.service.generate(export_id)
        status = self.service.status(export_id, "u1", "c1")
        self.assertEqual("completed", status["status"])
        self.assertEqual(1, status["row_count"])
        self.assertNotIn("stored_name", status)
        send_path, file_name = self.service.file_for_send(export_id, "u1", "c1")
        self.assertEqual(path, send_path)
        self.assertTrue(file_name.endswith(".xlsx"))
        self.service.mark_sent(export_id, "u1", "c1")
        self.assertIsNotNone(self.exports.rows[export_id]["sent_at"])
        with self.assertRaises(PermissionError):
            self.service.status(export_id, "other", "c1")

    def test_pending_and_interrupted_running_exports_resume_after_restart(self):
        first = self.service.create(self._snapshot(), "u1", "c1")
        second = self.service.create(self._snapshot(), "u1", "c1")
        self.exports.rows[second].update(
            status="running", started_at=self.clock[0] - timedelta(seconds=600),
        )
        recovered = self.service.recover_pending()
        self.assertEqual({first, second}, set(recovered))
        self.assertEqual("completed", self.exports.rows[first]["status"])
        self.assertEqual("completed", self.exports.rows[second]["status"])
        self.assertEqual([], self.service.recover_pending())

    def test_cleanup_is_idempotent_and_download_marks_expired(self):
        first = self.service.create(self._snapshot(), "u1", "c1")
        path = self.service.generate(first)
        expired_at = self.clock[0] + timedelta(hours=25)
        self.assertEqual(1, self.service.cleanup(now=expired_at))
        self.assertFalse(path.exists())
        self.assertEqual("expired", self.exports.rows[first]["status"])
        self.assertEqual(0, self.service.cleanup(now=expired_at))

        self.clock[0] = datetime(2026, 1, 2, 12, 0, 0)
        second = self.service.create(self._snapshot(snapshot_id="t" * 32), "u1", "c1")
        second_path = self.service.generate(second)
        self.clock[0] += timedelta(hours=25)
        with self.assertRaisesRegex(PermissionError, "expired"):
            self.service.download(second, "u1", "c1")
        self.assertEqual("expired", self.exports.rows[second]["status"])
        self.assertFalse(second_path.exists())

    def test_cleanup_removes_failed_temporary_files(self):
        payload = _typed_payload()
        payload["items"][0]["rows"] *= 2
        service = self._service(max_rows=1)
        export_id = service.create(self._snapshot(payload), "u1", "c1")
        with self.assertRaises(ExportLimitError):
            service.generate(export_id)
        temporary = self.export_dir / f"{export_id}.leftover.tmp"
        temporary.write_bytes(b"partial")
        self.assertEqual(0, service.cleanup())
        self.assertFalse(temporary.exists())
        self.assertEqual(0, service.cleanup())

    def test_from_env_and_invalid_configuration(self):
        directory = Path(self.directory.name) / "from-env"
        environment = {
            "KB_EXPORT_DIR": str(directory),
            "KB_EXPORT_ENABLE": "false",
            "KB_EXPORT_MAX_ROWS": "123",
            "KB_EXPORT_TTL_HOURS": "2.5",
            "KB_EXPORT_MAX_FILE_MB": "3.5",
        }
        with patch.dict(os.environ, environment, clear=True):
            service = ExportService.from_env(lambda: _Connection())
        self.assertFalse(service._enabled)
        self.assertEqual(123, service._max_rows)
        self.assertEqual(2.5, service._ttl_hours)
        self.assertEqual(3.5, service._max_file_mb)

        invalid_values = (
            ("KB_EXPORT_ENABLE", "perhaps"),
            ("KB_EXPORT_MAX_ROWS", "0"),
            ("KB_EXPORT_MAX_ROWS", "1.5"),
            ("KB_EXPORT_TTL_HOURS", "-1"),
            ("KB_EXPORT_MAX_FILE_MB", "nan"),
        )
        for key, value in invalid_values:
            with self.subTest(key=key, value=value):
                bad = dict(environment)
                bad[key] = value
                with patch.dict(os.environ, bad, clear=True):
                    with self.assertRaises(ValueError):
                        ExportService.from_env(lambda: _Connection())
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(ValueError):
                ExportService.from_env(lambda: _Connection())


if __name__ == "__main__":
    unittest.main()
