# -*- coding: utf-8 -*-
"""小易冻结查询结果的 XLSX 导出。

导出任务只记录条件/口径摘要；数据始终来自 ``kb_qa_query_snapshot``，
生成文件时不会重跑任何业务查询。连接由调用方提供，仓储不负责提交，
Service 在状态迁移边界显式 commit/rollback。
"""

import contextlib
import hashlib
import json
import math
import os
import secrets
import unicodedata
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_ALLOWED_TOOLS = (
    "daily_sales", "channel_sales", "channel_yoy", "query_sql", "product_stats",
)


class ExportLimitError(ValueError):
    """冻结结果超过行数或文件大小上限。"""


def _decode_json(value):
    if isinstance(value, bytes):
        value = value.decode("utf-8")
    return json.loads(value) if isinstance(value, str) else value


def _json(value):
    return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))


class QuerySnapshotRepository:
    """查询快照仓储；连接应使用 PyMySQL ``DictCursor``。"""

    def __init__(self, connection):
        self._connection = connection

    def save(self, snapshot_id, msg_id, sender_uid, conversation_id, payload):
        sql = (
            "INSERT INTO `kb_qa_query_snapshot` "
            "(`snapshot_id`, `msg_id`, `sender_uid`, `conversation_id`, "
            "`payload_json`, `created_at`) VALUES (%s, %s, %s, %s, %s, %s)"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (
                snapshot_id, msg_id, sender_uid, conversation_id,
                _json(payload), datetime.now(),
            ))

    def get_for_owner(self, snapshot_id, sender_uid, conversation_id):
        sql = (
            "SELECT `payload_json` FROM `kb_qa_query_snapshot` "
            "WHERE `snapshot_id` = %s AND `sender_uid` = %s "
            "AND `conversation_id` = %s"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (snapshot_id, sender_uid, conversation_id))
            row = cursor.fetchone()
        return _decode_json(row["payload_json"]) if row else None

    def latest_for_owner(self, sender_uid, conversation_id):
        sql = (
            "SELECT `payload_json` FROM `kb_qa_query_snapshot` "
            "WHERE `sender_uid` = %s AND `conversation_id` = %s "
            "ORDER BY `created_at` DESC, `snapshot_id` DESC LIMIT 1"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (sender_uid, conversation_id))
            row = cursor.fetchone()
        return _decode_json(row["payload_json"]) if row else None

    def latest_id_for_owner(self, sender_uid, conversation_id):
        sql = (
            "SELECT `snapshot_id` FROM `kb_qa_query_snapshot` "
            "WHERE `sender_uid` = %s AND `conversation_id` = %s "
            "ORDER BY `created_at` DESC, `snapshot_id` DESC LIMIT 1"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (sender_uid, conversation_id))
            row = cursor.fetchone()
        return row["snapshot_id"] if row else None


class ExportRepository:
    """导出任务仓储；每次状态更新都限定合法的当前状态。"""

    _FIELDS = (
        "export_id", "snapshot_id", "requester_uid", "conversation_id",
        "status", "file_name", "stored_name", "row_count", "file_size",
        "file_sha256", "freshness_status", "query_summary_json", "error",
        "expires_at", "created_at", "started_at", "finished_at", "downloaded_at", "sent_at",
    )

    def __init__(self, connection):
        self._connection = connection

    def create(self, *, export_id, snapshot_id, requester_uid, conversation_id,
               file_name, stored_name, freshness_status, query_summary,
               expires_at, created_at):
        sql = (
            "INSERT INTO `kb_qa_export` "
            "(`export_id`, `snapshot_id`, `requester_uid`, `conversation_id`, "
            "`status`, `file_name`, `stored_name`, `freshness_status`, "
            "`query_summary_json`, `expires_at`, `created_at`) "
            "VALUES (%s, %s, %s, %s, 'pending', %s, %s, %s, %s, %s, %s)"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (
                export_id, snapshot_id, requester_uid, conversation_id,
                file_name, stored_name, freshness_status, _json(query_summary),
                expires_at, created_at,
            ))
        return export_id

    def get(self, export_id):
        sql = (
            f"SELECT {self._select_fields()} FROM `kb_qa_export` "
            "WHERE `export_id` = %s"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (export_id,))
            row = cursor.fetchone()
        return dict(row) if row else None

    def list_expired(self, now):
        sql = (
            f"SELECT {self._select_fields()} FROM `kb_qa_export` "
            "WHERE `expires_at` <= %s AND `status` IN "
            "('pending', 'running', 'completed', 'failed') "
            "ORDER BY `expires_at`, `export_id`"
        )
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, (now,))
            rows = cursor.fetchall()
        return [dict(row) for row in rows]

    def claim_pending(self, now, *, stale_seconds=300):
        """原子认领 pending 或租约过期的 running 任务，供重启恢复。"""
        from common.public_data.db import transaction

        stale_before = now - timedelta(seconds=stale_seconds)
        with transaction(self._connection):
            with contextlib.closing(self._connection.cursor()) as cursor:
                cursor.execute(
                    "SELECT `export_id` FROM `kb_qa_export` "
                    "WHERE `expires_at` > %s AND (`status` = 'pending' OR "
                    "(`status` = 'running' AND `started_at` <= %s)) "
                    "ORDER BY `created_at`, `export_id` LIMIT 1 FOR UPDATE SKIP LOCKED",
                    (now, stale_before),
                )
                row = cursor.fetchone()
                if not row:
                    return None
                cursor.execute(
                    "UPDATE `kb_qa_export` SET `status` = 'running', "
                    "`started_at` = %s, `error` = NULL WHERE `export_id` = %s",
                    (now, row["export_id"]),
                )
                if cursor.rowcount != 1:
                    raise RuntimeError("export claim failed")
                return row["export_id"]

    def mark_running(self, export_id):
        return self._execute_transition(
            "UPDATE `kb_qa_export` SET `status` = 'running', "
            "`started_at` = UTC_TIMESTAMP(6), `error` = NULL "
            "WHERE `export_id` = %s AND `status` = 'pending'",
            (export_id,),
        )

    def mark_completed(self, export_id, *, stored_name, row_count, file_size,
                       file_sha256, finished_at):
        return self._execute_transition(
            "UPDATE `kb_qa_export` SET `status` = 'completed', "
            "`stored_name` = %s, `row_count` = %s, `file_size` = %s, "
            "`file_sha256` = %s, `error` = NULL, `finished_at` = %s "
            "WHERE `export_id` = %s AND `status` = 'running'",
            (stored_name, row_count, file_size, file_sha256, finished_at, export_id),
        )

    def mark_failed(self, export_id, error, finished_at=None):
        finished_at = finished_at or datetime.now()
        return self._execute_transition(
            "UPDATE `kb_qa_export` SET `status` = 'failed', `error` = %s, "
            "`finished_at` = %s WHERE `export_id` = %s "
            "AND `status` IN ('pending', 'running')",
            (str(error)[:4000], finished_at, export_id),
        )

    def mark_downloaded(self, export_id, downloaded_at=None):
        downloaded_at = downloaded_at or datetime.now()
        return self._execute_transition(
            "UPDATE `kb_qa_export` SET `downloaded_at` = %s "
            "WHERE `export_id` = %s AND `status` = 'completed'",
            (downloaded_at, export_id),
        )

    def mark_sent(self, export_id, sent_at=None):
        sent_at = sent_at or datetime.now()
        return self._execute_transition(
            "UPDATE `kb_qa_export` SET `sent_at` = %s "
            "WHERE `export_id` = %s AND `status` = 'completed'",
            (sent_at, export_id),
        )

    def expire(self, export_id):
        return self._execute_transition(
            "UPDATE `kb_qa_export` SET `status` = 'expired' "
            "WHERE `export_id` = %s AND `status` IN "
            "('pending', 'running', 'completed', 'failed')",
            (export_id,),
        )

    @classmethod
    def _select_fields(cls):
        return ", ".join(f"`{name}`" for name in cls._FIELDS)

    def _execute_transition(self, sql, params):
        with contextlib.closing(self._connection.cursor()) as cursor:
            cursor.execute(sql, params)
            return cursor.rowcount == 1


def _as_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError(f"invalid export timestamp: {value!r}") from exc
    raise ValueError(f"invalid export timestamp: {value!r}")


def _redact_sql(value):
    """任务摘要可以保留筛选条件，但永不复制完整 SQL。"""
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            name = str(key)
            if name.lower() == "sql" or name.lower().endswith("_sql"):
                result[name] = "[已省略]"
            else:
                result[name] = _redact_sql(item)
        return result
    if isinstance(value, (list, tuple)):
        return [_redact_sql(item) for item in value]
    return value


def _export_file_name(payload, created_at):
    themes = {
        "daily_sales": "区域日销",
        "channel_sales": "渠道日销",
        "channel_yoy": "渠道同比",
        "product_stats": "商品统计",
        "query_sql": "自定义查询",
    }
    items = payload.get("items") if isinstance(payload, dict) else []
    tool_name = items[0].get("tool_name") if items else None
    theme = themes.get(tool_name, "查询结果")
    return f"小易{theme}_{created_at:%Y%m%d_%H%M%S}.xlsx"


def _snapshot_summary(payload, allowed_tools):
    if not isinstance(payload, dict) or payload.get("version") != 1:
        raise ValueError("snapshot payload must have version=1")
    items = payload.get("items")
    if not isinstance(items, list):
        raise ValueError("snapshot payload items must be a list")

    summaries = []
    freshness_values = []
    total_rows = 0
    for item in items:
        if not isinstance(item, dict):
            raise ValueError("snapshot item must be an object")
        tool_name = item.get("tool_name")
        if tool_name not in allowed_tools:
            raise PermissionError(f"tool is not exportable: {tool_name!r}")
        rows = item.get("rows")
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("snapshot item rows must be a list of objects")
        total_rows += len(rows)
        freshness = item.get("freshness_status")
        if freshness is not None:
            freshness_values.append(str(freshness))
        summaries.append({
            "tool_name": tool_name,
            "effective_params": _redact_sql(item.get("effective_params") or {}),
            "row_count": len(rows),
            "data_scope": item.get("data_scope"),
            "data_version": item.get("data_version"),
            "freshness_status": freshness,
        })

    freshness_status = None
    if freshness_values:
        freshness_status = (
            freshness_values[0]
            if all(value == freshness_values[0] for value in freshness_values)
            else "mixed"
        )
    return {
        "version": 1,
        "row_count": total_rows,
        "freshness_status": freshness_status,
        "items": summaries,
    }


_FIELD_METADATA = {
    "stat_date": ("业务日期", "", "统计业务日期"),
    "business_date": ("业务日期", "", "实际业务归属日期"),
    "trade_date": ("业务日期", "", "订单业务归属日期"),
    "scope": ("区域", "", "业务区域或统计范围"),
    "region": ("区域", "", "业务区域"),
    "channel": ("渠道", "", "业务渠道"),
    "channel_key": ("渠道编码", "", "渠道唯一标识"),
    "channel_label": ("渠道名称", "", "渠道显示名称"),
    "store_name": ("门店", "", "门店名称"),
    "shop_name": ("店铺", "", "ERP/平台店铺名称"),
    "responsible_person": ("责任人", "", "业务责任人"),
    "sales_amount": ("销售额", "元", "冻结结果中的销售金额"),
    "paid_amount": ("实付金额", "元", "订单实付金额"),
    "amount": ("金额", "元", "业务金额"),
    "monthly_target": ("月目标", "元", "月度目标金额"),
    "dod_rate": ("日环比", "%", "与上一业务日相比"),
    "wow_rate": ("周环比", "%", "与上一周相比"),
    "month_rate": ("月完成率", "%", "月累计与月目标之比"),
    "yoy_rate": ("同比", "%", "与上年同期相比"),
    "order_count": ("订单数", "单", "订单数量"),
    "product_count": ("商品数", "款", "商品数量"),
    "brand": ("品牌", "", "商品品牌"),
    "category": ("类别", "", "商品类别"),
    "record_id": ("记录编号", "", "来源记录标识"),
    "year_month": ("月份", "", "业务月份"),
    "metric": ("指标", "", "统计指标名称"),
    "value": ("数值", "", "指标当前值"),
    "previous_value": ("上期值", "", "指标对比期值"),
}


def _columns_for(item):
    rows = item["rows"]
    raw_columns = item.get("columns")
    if raw_columns is None:
        raw_columns = []
        seen = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.add(key)
                    raw_columns.append(key)

    if not isinstance(raw_columns, (list, tuple)):
        raise ValueError("snapshot item columns must be a list")
    columns = []
    seen = set()
    for raw in raw_columns:
        if isinstance(raw, str):
            raw = {"key": raw}
        if not isinstance(raw, dict) or not isinstance(raw.get("key"), str) or not raw["key"]:
            raise ValueError("each column must be a key or an object with a key")
        key = raw["key"]
        if key in seen:
            raise ValueError(f"duplicate export column: {key}")
        seen.add(key)
        default_label, default_unit, default_description = _FIELD_METADATA.get(
            key, (key, "", ""),
        )
        label = raw.get("label") or default_label
        unit = raw.get("unit") if raw.get("unit") is not None else default_unit
        description = raw.get("description") or default_description
        columns.append({
            "key": key,
            "label": str(label),
            "unit": str(unit),
            "description": str(description),
        })
    return columns


def _safe_text(value):
    text = value if isinstance(value, str) else str(value)
    # Excel 在开头空白之后仍可能识别公式；保留原文，只加文本前导符。
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    if len(text) > 32767:
        raise ExportLimitError("Excel cell text exceeds 32767 characters")
    return text


def _excel_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        # XLSX 不保存时区；保留冻结值的年月日时分秒与微秒。
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, date):
        return value
    if isinstance(value, (dict, list, tuple)):
        return _safe_text(_json(value))
    return _safe_text(value)


def _header_text(column):
    label = column["label"]
    return f"{label}（{column['unit']}）" if column["unit"] else label


def _display_width(value):
    value = _excel_value(value)
    text = "" if value is None else str(value)
    return sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in text)


def _append_header(sheet, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=_safe_text(value))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cells.append(cell)
    sheet.append(cells)


def _layout_data_sheet(sheet, columns, rows):
    count = len(columns)
    if not count:
        return
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(count)}{len(rows) + 1}"
    for index, column in enumerate(columns, 1):
        values = [_header_text(column)]
        values.extend(row.get(column["key"]) for row in rows[:200])
        width = min(60, max(10, max(_display_width(value) for value in values) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width


def _append_data_row(sheet, columns, row):
    cells = []
    for column in columns:
        value = _excel_value(row.get(column["key"]))
        if isinstance(value, datetime):
            cell = WriteOnlyCell(sheet, value=value)
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cells.append(cell)
        elif isinstance(value, date):
            cell = WriteOnlyCell(sheet, value=value)
            cell.number_format = "yyyy-mm-dd"
            cells.append(cell)
        else:
            cells.append(value)
    sheet.append(cells)


def _write_snapshot_xlsx(payload, destination, *, requester_uid, exported_at,
                         query_summary=None):
    """将冻结 payload 写到指定路径，返回数据总行数。

    这是纯文件生成函数，没有数据库连接或业务查询入口，便于独立测试。
    """
    summary = query_summary or _snapshot_summary(payload, DEFAULT_ALLOWED_TOOLS)
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "小易"
    workbook.properties.title = "小易查询结果"
    all_columns = []
    total_rows = 0

    for index, item in enumerate(payload["items"], 1):
        title = "数据" if index == 1 else f"数据_{index}"
        sheet = workbook.create_sheet(title=title)
        columns = _columns_for(item)
        rows = item["rows"]
        all_columns.append((title, item, columns))
        total_rows += len(rows)
        # write-only 模式的冻结窗格等视图属性必须写在首个单元格之前。
        _layout_data_sheet(sheet, columns, rows)
        if columns:
            _append_header(sheet, [_header_text(column) for column in columns])
        for row in rows:
            _append_data_row(sheet, columns, row)

    notes = workbook.create_sheet(title="导出说明")
    notes.freeze_panes = "A2"
    _append_header(notes, ["项目", "内容"])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 100
    notes.append(["导出人", _safe_text(requester_uid)])
    notes.append(["导出时间", _safe_text(exported_at.isoformat(sep=" "))])
    notes.append(["总行数", total_rows])
    notes.append(["freshness", _safe_text(summary.get("freshness_status") or "未提供")])
    summary_items = summary.get("items") or []
    for index, item in enumerate(payload["items"], 1):
        item_summary = summary_items[index - 1] if index <= len(summary_items) else {}
        prefix = f"结果 {index}"
        notes.append([f"{prefix} - 工具", _safe_text(item.get("tool_name") or "")])
        notes.append([
            f"{prefix} - 条件摘要",
            _safe_text(_json(item_summary.get("effective_params") or {})),
        ])
        notes.append([f"{prefix} - 口径/数据范围", _safe_text(item.get("data_scope") or "未提供")])
        notes.append([f"{prefix} - 来源", _safe_text(item.get("tool_name") or "")])
        notes.append([f"{prefix} - 数据截至", _safe_text(item.get("data_version") or "未提供")])
        notes.append([f"{prefix} - 行数", len(item["rows"])])
        notes.append([f"{prefix} - freshness", _safe_text(item.get("freshness_status") or "未提供")])

    fields = workbook.create_sheet(title="字段说明")
    fields.freeze_panes = "A2"
    fields.auto_filter.ref = "A1:E1"
    _append_header(fields, ["工作表", "字段", "中文名称", "单位", "说明"])
    for title, _item, columns in all_columns:
        for column in columns:
            fields.append([
                title, _safe_text(column["key"]), _safe_text(column["label"]),
                _safe_text(column["unit"]), _safe_text(column["description"]),
            ])
    for letter, width in zip("ABCDE", (18, 28, 28, 14, 70)):
        fields.column_dimensions[letter].width = width

    workbook.save(destination)
    return total_rows


def _sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _positive_int(value, name):
    if isinstance(value, bool):
        raise ValueError(f"{name} must be a positive integer")
    try:
        result = int(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(f"{name} must be a positive integer") from exc
    if result < 1 or str(value).strip() != str(result):
        raise ValueError(f"{name} must be a positive integer")
    return result


def _positive_number(value, name):
    if isinstance(value, bool):
        raise ValueError(f"{name} must be positive")
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(f"{name} must be positive") from exc
    if not math.isfinite(result) or result <= 0:
        raise ValueError(f"{name} must be positive")
    return result


def _env_enabled(value):
    if value is None or not value.strip():
        return False
    normalized = value.strip().lower()
    if normalized in ("1", "true", "yes", "on"):
        return True
    if normalized in ("0", "false", "no", "off"):
        return False
    raise ValueError("KB_EXPORT_ENABLE must be a boolean")


class ExportService:
    """导出任务服务。``authorizer`` 接收 ``(requester_uid, conversation_id)``。"""

    def __init__(self, connection_factory, export_dir, *, enabled=True,
                 max_rows=50000, ttl_hours=24, max_file_mb=20,
                 allowed_tools=None, now=None,
                 snapshot_repository_factory=QuerySnapshotRepository,
                 export_repository_factory=ExportRepository):
        self._connection_factory = connection_factory
        self._export_dir = Path(export_dir)
        self._enabled = bool(enabled)
        self._max_rows = _positive_int(max_rows, "max_rows")
        self._ttl_hours = _positive_number(ttl_hours, "ttl_hours")
        self._max_file_mb = _positive_number(max_file_mb, "max_file_mb")
        self._allowed_tools = frozenset(allowed_tools or DEFAULT_ALLOWED_TOOLS)
        if not self._allowed_tools or any(not isinstance(tool, str) or not tool for tool in self._allowed_tools):
            raise ValueError("allowed_tools must contain tool names")
        self._now = now or datetime.now
        self._snapshot_repository_factory = snapshot_repository_factory
        self._export_repository_factory = export_repository_factory
        self._export_dir.mkdir(parents=True, exist_ok=True)

    @property
    def enabled(self):
        return self._enabled

    @classmethod
    def from_env(cls, connection_factory, environ=None):
        env = os.environ if environ is None else environ
        export_dir = (env.get("KB_EXPORT_DIR") or "").strip()
        if not export_dir:
            raise ValueError("KB_EXPORT_DIR is required")
        return cls(
            connection_factory,
            export_dir,
            enabled=_env_enabled(env.get("KB_EXPORT_ENABLE")),
            max_rows=_positive_int(env.get("KB_EXPORT_MAX_ROWS", "50000"), "KB_EXPORT_MAX_ROWS"),
            ttl_hours=_positive_number(env.get("KB_EXPORT_TTL_HOURS", "24"), "KB_EXPORT_TTL_HOURS"),
            max_file_mb=_positive_number(env.get("KB_EXPORT_MAX_FILE_MB", "20"), "KB_EXPORT_MAX_FILE_MB"),
        )

    def create(self, snapshot_id, requester_uid, conversation_id, authorizer=None):
        if not self._enabled:
            raise PermissionError("XLSX export is disabled")
        export_id = uuid.uuid4().hex
        created_at = self._now()
        connection = self._connection_factory()
        try:
            snapshots = self._snapshot_repository_factory(connection)
            exports = self._export_repository_factory(connection)
            payload = snapshots.get_for_owner(snapshot_id, requester_uid, conversation_id)
            if payload is None:
                raise PermissionError("snapshot is not owned by requester")
            self._authorize(authorizer, requester_uid, conversation_id)
            summary = _snapshot_summary(payload, self._allowed_tools)
            stored_name = f"{export_id}_{secrets.token_hex(8)}.xlsx"
            file_name = _export_file_name(payload, created_at)
            exports.create(
                export_id=export_id,
                snapshot_id=snapshot_id,
                requester_uid=requester_uid,
                conversation_id=conversation_id,
                file_name=file_name,
                stored_name=stored_name,
                freshness_status=summary["freshness_status"],
                query_summary=summary,
                expires_at=created_at + timedelta(hours=self._ttl_hours),
                created_at=created_at,
            )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        return export_id

    def latest_snapshot_id(self, requester_uid, conversation_id):
        connection = self._connection_factory()
        try:
            snapshots = self._snapshot_repository_factory(connection)
            return snapshots.latest_id_for_owner(requester_uid, conversation_id)
        finally:
            connection.close()

    def status(self, export_id, requester_uid, conversation_id, authorizer=None):
        """返回可安全透给请求人的任务状态，不含服务端存储名。"""
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            snapshots = self._snapshot_repository_factory(connection)
            task = self._require_task(exports, export_id)
            self._require_owner(task, snapshots, requester_uid, conversation_id)
            self._authorize(authorizer, requester_uid, conversation_id)
            if (task["status"] != "expired"
                    and _as_datetime(task["expires_at"]) <= self._now()):
                exports.expire(export_id)
                connection.commit()
                self._remove_task_file(task)
                task["status"] = "expired"
            return {
                "export_id": export_id,
                "snapshot_id": task["snapshot_id"],
                "status": task["status"],
                "file_name": task["file_name"],
                "row_count": task["row_count"],
                "file_size": task["file_size"],
                "freshness_status": task["freshness_status"],
                "query_summary": _decode_json(task["query_summary_json"]) or {},
                "expires_at": _as_datetime(task["expires_at"]).isoformat(sep=" "),
                "created_at": _as_datetime(task["created_at"]).isoformat(sep=" "),
                "finished_at": (
                    _as_datetime(task["finished_at"]).isoformat(sep=" ")
                    if task["finished_at"] else None
                ),
                "error": "导出失败，请重新发起" if task["status"] == "failed" else None,
            }
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def generate(self, export_id, authorizer=None, *, claimed=False):
        if not self._enabled:
            raise PermissionError("XLSX export is disabled")
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            snapshots = self._snapshot_repository_factory(connection)
            task = self._require_task(exports, export_id)
            requester_uid = task["requester_uid"]
            conversation_id = task["conversation_id"]
            payload = snapshots.get_for_owner(
                task["snapshot_id"], requester_uid, conversation_id,
            )
            if payload is None:
                raise PermissionError("snapshot is not owned by requester")
            self._authorize(authorizer, requester_uid, conversation_id)
            if _as_datetime(task["expires_at"]) <= self._now():
                exports.expire(export_id)
                connection.commit()
                self._remove_task_file(task)
                raise PermissionError("export has expired")
            if claimed:
                if task["status"] != "running":
                    raise RuntimeError("claimed export is no longer running")
            elif task["status"] != "pending" or not exports.mark_running(export_id):
                raise RuntimeError("only a pending export can be generated")
            # 冻结数据与摘要均在运行状态落库前读取，后续不再访问业务表。
            summary = _snapshot_summary(payload, self._allowed_tools)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

        temporary = None
        final = None
        try:
            if summary["row_count"] > self._max_rows:
                raise ExportLimitError(
                    f"export row count {summary['row_count']} exceeds limit {self._max_rows}"
                )
            final = self._stored_path(task["stored_name"])
            temporary = self._export_dir / f"{export_id}.{secrets.token_hex(8)}.tmp"
            row_count = _write_snapshot_xlsx(
                payload,
                temporary,
                requester_uid=requester_uid,
                exported_at=self._now(),
                query_summary=summary,
            )
            file_size = temporary.stat().st_size
            max_bytes = int(self._max_file_mb * 1024 * 1024)
            if file_size > max_bytes:
                raise ExportLimitError(
                    f"export file size {file_size} exceeds limit {max_bytes}"
                )
            os.replace(temporary, final)
            temporary = None
            file_sha256 = _sha256(final)

            connection = self._connection_factory()
            try:
                exports = self._export_repository_factory(connection)
                completed = exports.mark_completed(
                    export_id,
                    stored_name=task["stored_name"],
                    row_count=row_count,
                    file_size=file_size,
                    file_sha256=file_sha256,
                    finished_at=self._now(),
                )
                if not completed:
                    raise RuntimeError("running export no longer exists")
                connection.commit()
            except Exception:
                connection.rollback()
                raise
            finally:
                connection.close()
            return final
        except Exception as exc:
            for path in (temporary, final):
                self._unlink(path)
            self._mark_failed(export_id, exc)
            raise

    def download(self, export_id, requester_uid, conversation_id, authorizer=None):
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            snapshots = self._snapshot_repository_factory(connection)
            task = self._require_task(exports, export_id)
            self._require_owner(task, snapshots, requester_uid, conversation_id)
            self._authorize(authorizer, requester_uid, conversation_id)
            self._require_completed_file(task, exports, export_id, connection)
            if not exports.mark_downloaded(export_id, self._now()):
                raise RuntimeError("export is no longer completed")
            connection.commit()
            return self._stored_path(task["stored_name"]), task["file_name"]
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def file_for_send(self, export_id, requester_uid, conversation_id, authorizer=None):
        """群发送前重新校验归属、权限、有效期和文件完整性。"""
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            snapshots = self._snapshot_repository_factory(connection)
            task = self._require_task(exports, export_id)
            self._require_owner(task, snapshots, requester_uid, conversation_id)
            self._authorize(authorizer, requester_uid, conversation_id)
            self._require_completed_file(task, exports, export_id, connection)
            connection.commit()
            return self._stored_path(task["stored_name"]), task["file_name"]
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def mark_sent(self, export_id, requester_uid, conversation_id, authorizer=None):
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            snapshots = self._snapshot_repository_factory(connection)
            task = self._require_task(exports, export_id)
            self._require_owner(task, snapshots, requester_uid, conversation_id)
            self._authorize(authorizer, requester_uid, conversation_id)
            if not exports.mark_sent(export_id, self._now()):
                raise RuntimeError("export is no longer completed")
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def recover_pending(self, authorizer=None, *, limit=10):
        """进程重启后恢复未完成任务；单个失败不阻断其他任务。"""
        recovered = []
        for _ in range(max(1, min(int(limit), 100))):
            connection = self._connection_factory()
            try:
                exports = self._export_repository_factory(connection)
                export_id = exports.claim_pending(self._now())
            finally:
                connection.close()
            if export_id is None:
                break
            try:
                self.generate(export_id, authorizer, claimed=True)
                recovered.append(export_id)
            except Exception:
                continue
        return recovered

    def cleanup(self, now=None):
        current = now or self._now()
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            expired = exports.list_expired(current)
            for task in expired:
                self._remove_task_file(task)
                exports.expire(task["export_id"])

            # 正常失败路径会即时删 tmp；这里覆盖进程中断后已由其他流程标失败的残留。
            for temporary in self._export_dir.glob("*.tmp"):
                export_id = temporary.name.split(".", 1)[0]
                task = exports.get(export_id)
                if task and task.get("status") in ("failed", "expired"):
                    self._unlink(temporary)
            connection.commit()
            return len(expired)
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    @staticmethod
    def _require_task(exports, export_id):
        task = exports.get(export_id)
        if task is None:
            raise KeyError(f"unknown export_id: {export_id}")
        return task

    @staticmethod
    def _require_owner(task, snapshots, requester_uid, conversation_id):
        if (task["requester_uid"] != requester_uid
                or task["conversation_id"] != conversation_id):
            raise PermissionError("export is not owned by requester")
        payload = snapshots.get_for_owner(
            task["snapshot_id"], requester_uid, conversation_id,
        )
        if payload is None:
            raise PermissionError("snapshot is not owned by requester")

    def _require_completed_file(self, task, exports, export_id, connection):
        if _as_datetime(task["expires_at"]) <= self._now():
            exports.expire(export_id)
            connection.commit()
            self._remove_task_file(task)
            raise PermissionError("export has expired")
        if task["status"] != "completed":
            raise RuntimeError("export is not completed")
        path = self._stored_path(task["stored_name"])
        if not path.is_file():
            raise FileNotFoundError(path)

    @staticmethod
    def _authorize(authorizer, requester_uid, conversation_id):
        if authorizer is not None and authorizer(requester_uid, conversation_id) is False:
            raise PermissionError("export authorization denied")

    def _stored_path(self, stored_name):
        if not stored_name or Path(stored_name).name != stored_name:
            raise ValueError("invalid stored file name")
        return self._export_dir / stored_name

    def _remove_task_file(self, task):
        if task.get("stored_name"):
            self._unlink(self._stored_path(task["stored_name"]))

    @staticmethod
    def _unlink(path):
        if path is not None:
            try:
                Path(path).unlink()
            except FileNotFoundError:
                pass

    def _mark_failed(self, export_id, error):
        connection = self._connection_factory()
        try:
            exports = self._export_repository_factory(connection)
            exports.mark_failed(export_id, str(error) or type(error).__name__, self._now())
            connection.commit()
        except Exception:
            connection.rollback()
        finally:
            connection.close()
